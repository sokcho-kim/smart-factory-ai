"""기획서 최종본 — 텍스트박스 크기 보정 + 내용 간결화."""
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC_TEMPLATE = "C:/Users/sokch/Downloads/2026 스마트공장 운영시스템 MVP 개발해커톤 기획서 양식.pptx"
DST = "docs/proposal/기획서_final2.pptx"

prs = Presentation(SRC_TEMPLATE)
slides = list(prs.slides)

# 내용 영역 좌표 (배경 박스 안, 구분선 아래)
CONTENT_LEFT = Inches(1.0)
CONTENT_TOP = Inches(1.85)
CONTENT_WIDTH = Inches(10.8)
CONTENT_HEIGHT = Inches(4.4)


def find_and_resize_content(slide):
    """TextBox 10 (작성방법 박스)을 찾아서 크기 조정 후 반환."""
    for shape in slide.shapes:
        if shape.has_text_frame and "작성방법" in shape.text_frame.text:
            shape.left = CONTENT_LEFT
            shape.top = CONTENT_TOP
            shape.width = CONTENT_WIDTH
            shape.height = CONTENT_HEIGHT
            return shape
    return None


def write_lines(shape, lines, font_size=10):
    """텍스트 프레임에 줄 단위로 작성."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        run = p.runs[0] if p.runs else p.add_run()
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        if line.startswith("[") and "]" in line:
            run.font.bold = True
            run.font.size = Pt(font_size + 1)
            run.font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
        elif line == "":
            run.font.size = Pt(4)
        elif line.startswith("  "):
            run.font.size = Pt(font_size - 1)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        p.space_after = Pt(1)
        p.space_before = Pt(0)


# ═══ Slide 1: 표지 ═══
for shape in slides[0].shapes:
    if shape.has_table:
        table = shape.table
        table.cell(0, 1).text = "Fist99"
        table.cell(1, 1).text = "PunchLine - AI 비전 기반 품질 자동 기록 및 불량 원인 추적 시스템"
        for row in table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(12)

# ═══ Slide 2: 문제 정의 ═══
box = find_and_resize_content(slides[1])
if box:
    write_lines(box, [
        "[문제: 데이터가 안 쌓인다 -- 그래서 분석도 추적도 안 된다]",
        "",
        "중소 제조 현장의 품질 관리는 여전히 사람 눈과 수기 기록에 의존합니다.",
        "",
        "1) 검사 데이터가 부정확하거나 안 쌓인다",
        '  작업자가 바쁘면 기록을 건너뛰고, 시간을 대충 적습니다.',
        '  "몇 시에 불량이 급증했는가?"에 답할 수 없습니다. (JPC 구매부 인터뷰)',
        "",
        "2) 불량이 반복되는데 원인을 모른다",
        "  불량 시점의 공정 조건(온도/압력/속도)이 기록과 연결되지 않습니다.",
        "",
        "3) AI가 필요한 건 알지만 도입 못 한다",
        "  중소 제조기업 80%가 AI 필요성 체감 (중기중앙회, 2026)",
        "  그러나 AI 비전 도입의 77%가 파일럿에서 멈춤 (ABI Research)",
        "  원인: 산업용 카메라/조명/GPU 서버 등 초기 투자 비용",
        "",
        "[핵심 KPI]",
        "  품질 데이터 수집률:  0% --> 100% (사람 입력 제거)",
        "  불량 원인 추적:     수일(수작업) --> 실시간",
        "  도입 비용:          수천만원 --> 웹캠 + 일반 PC",
    ])

# ═══ Slide 3: 솔루션 개요 ═══
box = find_and_resize_content(slides[2])
if box:
    write_lines(box, [
        "[PunchLine -- 카메라가 검사하고, 시간을 기록하고, 원인을 찾는다]",
        "",
        "사람이 입력하는 것이 하나도 없는 품질 관리 시스템입니다.",
        "라인에 카메라를 설치하면, AI가 제품을 검사하고,",
        "불량을 자동 기록하고, 공정 조건과 연결하여 원인을 추적합니다.",
        "",
        "[핵심: 분석 도구가 아니라 '데이터 생성 엔진']",
        "  기존 솔루션은 분석 도구를 제공하지만, 현장에는 분석할 데이터가 없습니다.",
        "  PunchLine은 카메라가 촬영하는 순간 검사결과+시간+유형이 자동 기록됩니다.",
        "",
        "[2026 제조 AI 트렌드 부합]",
        "  Predictive Quality  -- 탐지를 넘어 예측/예방",
        "  Causal AI           -- 상관관계가 아닌 진짜 원인 추적",
        "  저비용 진입          -- 웹캠 + YOLO v8n으로 시작 가능",
        "",
        "[핵심 흐름]",
        "  촬영 --> DETECT(불량판정) --> TRACE(원인추적) --> GUIDE(조치안내) --> 자동기록",
        "  * 전 과정 사람 입력 0건",
    ])

# ═══ Slide 4: 주요 기능 정의 ═══
box = find_and_resize_content(slides[3])
if box:
    write_lines(box, [
        "[기능 1. 저비용 비전 자동 검사 + 자동 기록]",
        "  웹캠 또는 저가 USB 카메라로 제품 외관 촬영",
        "  경량 모델(YOLO v8n)이 양품/불량 판정 + 유형 분류",
        "  연기/진동/소음 환경에서도 동작 검증 완료 (팀원 실전 경험)",
        "  검사 결과 + 타임스탬프 자동 DB 적재",
        "",
        "[기능 2. 정확한 시계열 품질 추적]",
        '  "몇 시에 불량이 급증했는가?" 즉시 확인',
        "  시간대별/유형별/라인별 불량률 추이 자동 생성",
        "  교대조별 품질 비교, 요일/시간 패턴 분석",
        "",
        "[기능 3. 불량 원인 자동 추적 (인과 분석)]",
        "  불량 급증 시점의 공정 파라미터 자동 조회",
        "  다변수 복합 상호작용 분석 (온도+압력+속도)",
        "  SHAP 기반 설명가능 AI로 핵심 원인 변수 특정",
        "",
        "[기능 4. 조치 가이드 + 이력 관리]",
        "  불량 유형별 표준 조치 방법 자동 안내",
        "  과거 유사 사례 및 해결 이력 검색",
        "  전체 품질 이력 자동 축적 -- 감사/인증 즉시 대응",
    ], font_size=9)

# ═══ Slide 5: 데이터 및 기술 활용 계획 ═══
box = find_and_resize_content(slides[4])
if box:
    write_lines(box, [
        "[활용 데이터]",
        "  비전 학습:    MVTec AD (5,354장) / Casting Defect (7,348장) / NEU (1,800장)",
        "  인과 분석:    SECOM 반도체 센서 590개 + 불량라벨 (1,567 rows)",
        "  시뮬레이션:   온도/압력/속도 복합 상호작용 합성 데이터",
        "  * 제품 도면(CAD) 불필요 -- 양품 이미지만으로 학습 가능",
        "",
        "[기술 구성]",
        "  결함 탐지:    YOLO v8n (경량, GPU 없이 추론 가능) + PatchCore (이상 탐지)",
        "  인과 분석:    SHAP + 다변수 상호작용 분석",
        "  Backend:     FastAPI + PostgreSQL",
        "  Frontend:    Streamlit + Plotly",
        "  Infra:       Docker Compose (FE+BE+DB 일체화)",
        "",
        "[하드웨어 -- 최소 사양으로 동작]",
        "  카메라:       웹캠 또는 저가 USB 카메라 (3~10만원)",
        "  PC:          일반 사무용 PC (GPU 불필요)",
        "  네트워크:     폐쇄망 설치 가능 (Docker, 외부 통신 없음)",
        "",
        "[정책 연계]",
        "  2026 스마트 제조혁신 AI트랙 (154.5억원, 150개 중소기업 대상)",
        "  기존 스마트팩토리 --> AI 팩토리 업그레이드 정책 부합",
    ], font_size=9)

# ═══ Slide 6: 사용자 시나리오 ═══
box = find_and_resize_content(slides[5])
if box:
    write_lines(box, [
        "[사용자]",
        "  품질 관리자 -- 불량 모니터링, 원인 파악, 개선 지시",
        "  라인 작업자 -- 불량 알림 수신, 현장 조치",
        "  공장장     -- 일일/주간 품질 리포트, 의사결정",
        "",
        '[기존 방식 -- 사람이 보고, 사람이 적는다]',
        "  09:47  작업자가 불량 발견, 수기 기록 (시간 부정확)",
        "  12:00  점심시간 기록 누락",
        "  15:00  품질 관리자가 엑셀 확인 -- 데이터 빠져있음",
        '  15:30  "오전에 불량 많았다는데 몇 시인지 모르겠다"',
        "  --> 원인 추적 불가, 내일도 반복",
        "",
        "[PunchLine 도입 후 -- 카메라가 보고, AI가 적는다]",
        "  09:00  라인 가동, 웹캠이 매 제품 자동 촬영",
        '  09:42  크랙 3건/10분 감지, 자동 알림 발생',
        '         "A라인 09:42부터 크랙 급증"',
        '         "금형 온도 86도 -- 정상 범위(70~80) 초과, 기여도 73%"',
        "  09:44  작업자 온도 조정",
        "  09:55  불량률 정상 복귀, 전 과정 자동 기록",
        "",
        "  감지 --> 원인 --> 조치: 2분",
    ], font_size=9)

# ═══ Slide 7: MVP 구현 범위 ═══
box = find_and_resize_content(slides[6])
if box:
    write_lines(box, [
        "[MVP 구현 범위 및 완성도 기준 (100점 만점)]",
        "",
        "  구현 항목                              배점   완성 기준",
        "  -------------------------------------------------------",
        "  AI 자동 검사 (양품/불량 판정)           25점   정확도 90%+, 실시간 판정 시연",
        "  불량 유형 분류 (크랙/기포/스크래치)      15점   3종 이상 분류, 바운딩박스 표시",
        "  자동 기록 + 시계열 추적                 20점   타임스탬프 DB 적재, 시간대별 차트",
        "  불량 원인 추적 (공정 조건 매칭)          20점   SHAP 분석, 원인 변수 Top3 표시",
        "  조치 가이드 제공                        10점   유형별 조치 안내, 유사 사례 검색",
        "  통합 대시보드 시연                      10점   전체 흐름 E2E 동작, Docker 배포",
        "  -------------------------------------------------------",
        "  목표: 85점 이상 (조치 가이드 일부 미구현 허용)",
        "",
        "[역할 분담]",
        "  원정환 -- AI 검사(25) + 유형분류(15) + 대시보드(10) = 50점 영역",
        "    YOLO v8n 학습+추론 / PatchCore 이상탐지 / Streamlit+Docker",
        "    (열악한 환경 실전 검증 경험 보유: 연기/진동/소음 속 YOLO 동작)",
        "",
        "  소지민 -- 자동기록(20) + 원인추적(20) + 조치가이드(10) = 50점 영역",
        "    시계열 품질 추적 / 다변수 SHAP 인과분석 / 지식베이스 구축",
        "",
        "[사전 준비] 모델 학습, 분석 파이프라인, 시뮬레이션 데이터 -- 본선 전 완료",
        "[본선 당일] 통합 + E2E 시연 + 발표에 집중",
        "[향후 확장] 라인카메라 직결 / MES 연동 / 치수검사(CAD) / 다공장 관제",
    ], font_size=9)

# ═══ Slide 8: 기대 효과 ═══
box = find_and_resize_content(slides[7])
if box:
    write_lines(box, [
        "[1. 인력난 해결 (제조업 AI 도입 기대효과 1위, 40%)]",
        "  육안 검사 --> AI 자동 검사 (탐지율 70% --> 95%)",
        "  숙련자 은퇴해도 품질 유지, 3교대 검사 --> 1인 모니터링",
        "",
        "[2. 데이터가 알아서 쌓인다]",
        "  수기 --> 카메라 자동 기록 (수집률 0% --> 100%)",
        "  정확한 타임스탬프로 시계열 분석 가능",
        "  감사/인증 시 품질 이력 즉시 제출",
        "",
        "[3. 불량 원인을 알 수 있다]",
        "  불량 시점 + 공정 조건 자동 매칭으로 인과 분석",
        "  반복 불량 50% 감소 목표",
        "",
        "[4. 이빨이 없어도 잇몸으로 시작할 수 있다]",
        "  웹캠(3만원) + 일반 PC면 도입 가능",
        "  양품 이미지 100장이면 학습 시작 (제품 도면 불필요)",
        "  폐쇄망 Docker 설치, 보안 걱정 없음",
        "  AI 비전 77%가 파일럿에서 멈추는 이유(비용)를 제거",
        "",
        "[정책 연계]  2026 스마트 제조혁신 AI트랙 (154.5억, 150개사)",
        "[확장 방향]  치수 검사 --> MES 연동 --> 다공장 관제 --> 디지털 트윈",
    ], font_size=9)

prs.save(DST)
print("저장 완료:", DST)
